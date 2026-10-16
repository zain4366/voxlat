"""Build the VoxLat manuscript .docx in the layout of the attached TAR template.

Source: meta.py (front matter), body_*.txt (lightweight markup), refs.py (references).
Equations: LaTeX -> OMML through pandoc (batch), inserted into python-docx paragraphs.
"""
import copy
import re
import subprocess
import sys
import tempfile
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from lxml import etree

import meta
from refs import REFS

HERE = Path(__file__).parent
TEMPLATE = HERE / "template_base.docx"
REVIEW = "--review" in sys.argv
OUT = HERE / "out" / ("VoxLat_manuscript_review.docx" if REVIEW else "VoxLat_manuscript_template_layout.docx")
LINE = 1.5 if REVIEW else None
BODY = ["body_1.txt", "body_2.txt", "body_3.txt", "body_4.txt"]

W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
M_NS = "http://schemas.openxmlformats.org/officeDocument/2006/math"
TEXT_W_TW = 9746                     # text width in twips (A4, 1.905 cm margins)
COL_W_TW = (TEXT_W_TW - 400) // 2    # two columns, 400 twips apart
BODY_PT = 10
SMALL_PT = 8

# ----------------------------------------------------------------------------- source parsing


def parse_source(text):
    blocks, para = [], []
    lines = text.split("\n")
    i = 0

    def flush():
        if para:
            blocks.append({"type": "para", "text": " ".join(s.strip() for s in para)})
            para.clear()

    while i < len(lines):
        ln = lines[i]
        s = ln.strip()
        if not s:
            flush(); i += 1; continue
        m = re.match(r"^(#{1,3})(\*?)\s+(.*?)\s*(\{#(sec:\w+)\})?$", s)
        if m:
            flush()
            blocks.append({"type": f"h{len(m.group(1))}", "text": m.group(3), "label": m.group(5),
                           "numbered": m.group(2) != "*"})
            i += 1; continue
        m = re.match(r"^\$\$(.*)\$\$\s*(\{#(eq:\w+)\})?$", s)
        if m:
            flush(); blocks.append({"type": "eq", "tex": m.group(1).strip(), "label": m.group(3)}); i += 1; continue
        m = re.match(r"^!\[(fig:\w+)\|(\w+)(?:\|([\d.]+))?\]\((.*?)\)\s*(.*)$", s)
        if m:
            flush()
            blocks.append({"type": "fig", "label": m.group(1), "width": m.group(2),
                           "cm": float(m.group(3)) if m.group(3) else None, "path": m.group(4), "caption": m.group(5)})
            i += 1; continue
        m = re.match(r"^TABLE\s+(tbl:\w+)\|(\w+)(?:\|([\d,]+))?$", s)
        if m:
            flush()
            tb = {"type": "table", "label": m.group(1), "width": m.group(2),
                  "cols": [int(x) for x in m.group(3).split(",")] if m.group(3) else None, "rows": [], "caption": "",
                  "note": None}
            i += 1
            while lines[i].strip() != "ENDTABLE":
                t = lines[i].strip()
                if t.startswith("CAPTION:"):
                    tb["caption"] = t[len("CAPTION:"):].strip()
                elif t.startswith("|"):
                    tb["rows"].append([c.strip() for c in t.strip("|").split(" | ")])
                i += 1
            i += 1
            if i < len(lines) and lines[i].strip().startswith("NOTE:"):
                tb["note"] = lines[i].strip()[5:].strip(); i += 1
            blocks.append(tb); continue
        m = re.match(r"^(- |[a-z]\) |\d+\. )(.*)$", s)
        if m and not para:
            flush()
            marker = m.group(1).strip()
            blocks.append({"type": "item", "marker": "•" if marker == "-" else marker, "text": m.group(2)})
            i += 1; continue
        para.append(ln); i += 1
    flush()
    return blocks


# ----------------------------------------------------------------------------- numbering


def number_labels(blocks):
    lab = {}
    h1 = h2 = eq = fig = tbl = 0
    for b in blocks:
        if b["type"] == "h1" and b["numbered"]:
            h1 += 1; h2 = 0; b["num"] = f"{h1}."
            if b["label"]: lab[b["label"]] = f"{h1}"
        elif b["type"] == "h2":
            h2 += 1; b["num"] = f"{h1}.{h2}."
            if b["label"]: lab[b["label"]] = f"{h1}.{h2}"
        elif b["type"] == "eq":
            eq += 1; b["num"] = eq
            if b["label"]: lab[b["label"]] = str(eq)
        elif b["type"] == "fig":
            fig += 1; b["num"] = fig; lab[b["label"]] = str(fig)
        elif b["type"] == "table":
            tbl += 1; b["num"] = tbl; lab[b["label"]] = str(tbl)
    return lab


class Citer:
    def __init__(self):
        self.order = []

    def num(self, key):
        if key not in REFS:
            raise KeyError(f"missing reference {key}")
        if key not in self.order:
            self.order.append(key)
        return self.order.index(key) + 1

    def fmt(self, keys):
        ns = sorted(self.num(k) for k in keys)
        parts, i = [], 0
        while i < len(ns):
            j = i
            while j + 1 < len(ns) and ns[j + 1] == ns[j] + 1:
                j += 1
            parts.append(f"{ns[i]}–{ns[j]}" if j - i >= 2 else (f"{ns[i]}, {ns[j]}" if j == i + 1 else f"{ns[i]}"))
            i = j + 1
        return "[" + ", ".join(parts) + "]"


def resolve(text, lab, cit):
    def ref(m):
        kind, key = m.group(1), m.group(0)[1:]
        if key not in lab:
            raise KeyError(f"unknown label {key}")
        n = lab[key]
        return {"fig": f"Fig. {n}", "tbl": f"Table {n}", "sec": f"Section {n}", "eq": n}[kind]
    text = re.sub(r"\[(@[\w]+(?:;\s*@[\w]+)*)\]", lambda m: cit.fmt([k.strip()[1:] for k in m.group(1).split(";")]), text)
    text = re.sub(r"@(fig|tbl|sec|eq):\w+", ref, text)
    text = re.sub(r"Section ([\d.]+), Section ([\d.]+)", r"Sections \1 and \2", text)
    return text


# ----------------------------------------------------------------------------- math (pandoc)


class MathCache:
    def __init__(self):
        self.items = {}

    def collect(self, tex, display):
        self.items.setdefault((tex, display), None)

    def convert(self):
        keys = list(self.items)
        md = []
        for k, (tex, disp) in enumerate(keys):
            md.append(f"MARK{k}END")
            md.append(f"$${tex}$$" if disp else f"${tex}$")
        with tempfile.TemporaryDirectory() as td:
            src = Path(td) / "m.md"; out = Path(td) / "m.docx"
            src.write_text("\n\n".join(md), encoding="utf-8")
            subprocess.run(["pandoc", str(src), "-f", "markdown+tex_math_dollars", "-o", str(out)], check=True)
            xml = zipfile.ZipFile(out).read("word/document.xml")
        root = etree.fromstring(xml)
        body = root.find(f"{{{W_NS}}}body")
        paras = [p for p in body if p.tag == f"{{{W_NS}}}p"]
        cur = None
        for p in paras:
            txt = "".join(t.text or "" for t in p.iter(f"{{{W_NS}}}t"))
            m = re.fullmatch(r"MARK(\d+)END", txt.strip())
            if m:
                cur = int(m.group(1)); continue
            if cur is None:
                continue
            om = p.find(f".//{{{M_NS}}}oMath")
            if om is None:
                raise ValueError(f"pandoc failed on: {keys[cur][0]}")
            self.items[keys[cur]] = om
            cur = None
        missing = [k for k, v in self.items.items() if v is None]
        if missing:
            raise ValueError(f"no OMML for {missing[:3]}")

    def get(self, tex, display):
        return copy.deepcopy(self.items[(tex, display)])


# ----------------------------------------------------------------------------- docx helpers


def set_run_font(r, size=None, bold=None, italic=None, sup=False, sub=False, mono=False, color=None):
    if size: r.font.size = Pt(size)
    if bold is not None: r.bold = bold
    if italic is not None: r.italic = italic
    if sup: r.font.superscript = True
    if sub: r.font.subscript = True
    if mono: r.font.name = "Consolas"
    if color: r.font.color.rgb = RGBColor.from_string(color)


TOKEN = re.compile(r"(\$[^$]+\$)|(\*\*.+?\*\*)|(`[^`]+`)|((?<![\w*])\*(?=\S)[^*\n]+?(?<=\S)\*(?![\w*]))")
SUBRULE = re.compile(r"([A-Za-zρσδΔφηκμτθ])_(\{[^}]+\}|[A-Za-z0-9]+)")


def add_text(p, text, size, bold=False, italic=False, mathc=None, subs=False):
    pos = 0
    for m in TOKEN.finditer(text):
        if m.start() > pos:
            add_plain(p, text[pos:m.start()], size, bold, italic, subs)
        tok = m.group(0)
        if m.group(1):
            p._p.append(mathc.get(tok[1:-1], False))
        elif m.group(2):
            add_text(p, tok[2:-2], size, True, italic, mathc, subs)
        elif m.group(3):
            r = p.add_run(tok[1:-1]); set_run_font(r, size=size - 1, bold=bold, mono=True)
        else:
            add_text(p, tok[1:-1], size, bold, True, mathc, subs)
        pos = m.end()
    if pos < len(text):
        add_plain(p, text[pos:], size, bold, italic, subs)


def add_plain(p, text, size, bold, italic, subs):
    if not subs:
        r = p.add_run(text); set_run_font(r, size=size, bold=bold, italic=italic); return
    pos = 0
    for m in SUBRULE.finditer(text):
        r = p.add_run(text[pos:m.start()] + m.group(1)); set_run_font(r, size=size, bold=bold, italic=italic)
        sub = m.group(2).strip("{}")
        r = p.add_run(sub); set_run_font(r, size=size, bold=bold, italic=italic, sub=True)
        pos = m.end()
    r = p.add_run(text[pos:]); set_run_font(r, size=size, bold=bold, italic=italic)


def collect_math(text, mathc):
    for m in TOKEN.finditer(text):
        if m.group(1):
            mathc.collect(m.group(1)[1:-1], False)
        elif m.group(2):
            collect_math(m.group(2)[2:-2], mathc)
        elif m.group(4):
            collect_math(m.group(4)[1:-1], mathc)


def pfmt(p, align=None, before=0, after=5, first=None, left=None, hanging=None, keep_next=False, line=None):
    f = p.paragraph_format
    f.space_before = Pt(before); f.space_after = Pt(after)
    if align is not None: p.alignment = align
    if first is not None: f.first_line_indent = Cm(first)
    if left is not None: f.left_indent = Cm(left)
    if hanging is not None: f.first_line_indent = Cm(-hanging)
    if keep_next: f.keep_with_next = True
    if line: f.line_spacing = line


def outline(p, lvl):
    e = OxmlElement("w:outlineLvl"); e.set(qn("w:val"), str(lvl)); p._p.get_or_add_pPr().append(e)


def fld(p, kind):
    r = p.add_run(); e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), kind); r._r.append(e)


def instr(p, text):
    r = p.add_run(); e = OxmlElement("w:instrText"); e.set(qn("xml:space"), "preserve"); e.text = text; r._r.append(e)


TOC_ENTRIES, TOC_PAGES = [], {}
BREAK_BEFORE = set()


def cell_borders(cell, **edges):
    tcPr = cell._tc.get_or_add_tcPr()
    b = tcPr.find(qn("w:tcBorders"))
    if b is None:
        b = OxmlElement("w:tcBorders"); tcPr.append(b)
    for edge, sz in edges.items():
        e = OxmlElement(f"w:{edge}")
        if sz:
            e.set(qn("w:val"), "single"); e.set(qn("w:sz"), str(sz)); e.set(qn("w:color"), "000000")
        else:
            e.set(qn("w:val"), "nil")
        e.set(qn("w:space"), "0"); b.append(e)


def cell_margins(cell, top=30, bottom=30, left=50, right=50):
    tcPr = cell._tc.get_or_add_tcPr()
    mar = OxmlElement("w:tcMar")
    for k, v in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        e = OxmlElement(f"w:{k}"); e.set(qn("w:w"), str(v)); e.set(qn("w:type"), "dxa"); mar.append(e)
    tcPr.append(mar)


def table_setup(tbl, widths_tw, total_tw):
    tblPr = tbl._tbl.tblPr
    for tag in ("w:tblW", "w:tblLayout", "w:tblBorders"):
        old = tblPr.find(qn(tag))
        if old is not None: tblPr.remove(old)
    w = OxmlElement("w:tblW"); w.set(qn("w:w"), str(total_tw)); w.set(qn("w:type"), "dxa"); tblPr.append(w)
    lay = OxmlElement("w:tblLayout"); lay.set(qn("w:type"), "fixed"); tblPr.append(lay)
    bd = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{edge}"); e.set(qn("w:val"), "nil"); bd.append(e)
    tblPr.append(bd)
    grid = tbl._tbl.tblGrid
    for gc, wt in zip(grid.findall(qn("w:gridCol")), widths_tw):
        gc.set(qn("w:w"), str(wt))
    for row in tbl.rows:
        for c, wt in zip(row.cells, widths_tw):
            c.width = Cm(wt / 567.0)
            tcW = c._tc.get_or_add_tcPr().find(qn("w:tcW"))
            tcW.set(qn("w:w"), str(wt)); tcW.set(qn("w:type"), "dxa")


def no_split_row(row):
    trPr = row._tr.get_or_add_trPr()
    e = OxmlElement("w:cantSplit"); trPr.append(e)


def header_row(row):
    trPr = row._tr.get_or_add_trPr()
    e = OxmlElement("w:tblHeader"); trPr.append(e)


# ----------------------------------------------------------------------------- builder


class Builder:
    def __init__(self):
        self.doc = Document(str(TEMPLATE))
        body = self.doc.element.body
        self.final_sect = body.find(qn("w:sectPr"))
        first = None
        for el in list(body):
            if el.tag == qn("w:p") and el.find(f".//{qn('w:sectPr')}") is not None and first is None:
                first = copy.deepcopy(el.find(f".//{qn('w:sectPr')}"))
            if el is not self.final_sect:
                body.remove(el)
        self.first_sect = first          # carries the footer reference
        self.mode = 1
        self.nsect = 0
        self.after_heading = True

    # sections ---------------------------------------------------------------
    def _sect(self, cols, first):
        s = copy.deepcopy(self.first_sect if first else self.final_sect)
        if not first:
            t = s.find(qn("w:type"))
            if t is None:
                t = OxmlElement("w:type"); s.insert(0, t)
            t.set(qn("w:val"), "continuous")
        c = s.find(qn("w:cols"))
        for k in list(c.attrib): del c.attrib[k]
        if cols == 2:
            c.set(qn("w:num"), "2")
            c.set(qn("w:space"), "400")
        else:
            c.set(qn("w:space"), "720")
        if REVIEW:
            ln = OxmlElement("w:lnNumType")
            ln.set(qn("w:countBy"), "1"); ln.set(qn("w:distance"), "280"); ln.set(qn("w:restart"), "continuous")
            pg = s.find(qn("w:pgMar")); pg.addnext(ln)
        return s

    def mode_to(self, cols):
        if REVIEW:
            cols = 1
        if cols == self.mode:
            return
        p = self.doc.add_paragraph()
        pfmt(p, after=0)
        r = p.add_run(""); set_run_font(r, size=1)
        p._p.get_or_add_pPr().append(self._sect(self.mode, first=self.nsect == 0))
        self.nsect += 1
        self.mode = cols

    def finish(self):
        s = self._sect(self.mode, first=False)
        body = self.doc.element.body
        body.remove(self.final_sect)
        body.append(s)

    # paragraphs -------------------------------------------------------------
    def para(self, text, mathc, indent=True):
        p = self.doc.add_paragraph()
        pfmt(p, align=WD_ALIGN_PARAGRAPH.LEFT if "`" in text else WD_ALIGN_PARAGRAPH.JUSTIFY, after=4,
             first=0.0 if (self.after_heading or not indent) else 0.45, line=LINE)
        add_text(p, text, BODY_PT, mathc=mathc)
        self.after_heading = False
        return p

    def heading(self, b):
        p = self.doc.add_paragraph()
        if b["type"] == "h1":
            pfmt(p, before=10, after=5, keep_next=True)
            txt = (b.get("num", "") + " " + b["text"]).strip() if b["numbered"] else b["text"]
            add_text(p, txt, 11, bold=True)
            outline(p, 0)
        else:
            pfmt(p, before=7, after=3.5, keep_next=True)
            add_text(p, b["num"] + " " + b["text"], BODY_PT, bold=True)
            outline(p, 1)
        self.after_heading = True

    def item(self, b, mathc):
        p = self.doc.add_paragraph()
        pfmt(p, align=WD_ALIGN_PARAGRAPH.JUSTIFY, after=3, left=0.55, hanging=0.4, line=LINE)
        r = p.add_run(b["marker"] + "\t"); set_run_font(r, size=BODY_PT)
        tabs = p.paragraph_format.tab_stops; tabs.add_tab_stop(Cm(0.55))
        add_text(p, b["text"], BODY_PT, mathc=mathc)
        self.after_heading = False

    def equation(self, b, mathc):
        width = COL_W_TW if self.mode == 2 else TEXT_W_TW
        tbl = self.doc.add_table(rows=1, cols=2)
        table_setup(tbl, [width - 620, 620], width)
        c0, c1 = tbl.rows[0].cells
        for c in (c0, c1):
            cell_margins(c, 20, 20, 0, 0)
        p0 = c0.paragraphs[0]; pfmt(p0, align=WD_ALIGN_PARAGRAPH.CENTER, before=2, after=2)
        omp = OxmlElement("m:oMathPara")
        ompr = OxmlElement("m:oMathParaPr"); jc = OxmlElement("m:jc"); jc.set(qn("m:val"), "center"); ompr.append(jc)
        omp.append(ompr); omp.append(mathc.get(b["tex"], True)); p0._p.append(omp)
        p1 = c1.paragraphs[0]; pfmt(p1, align=WD_ALIGN_PARAGRAPH.RIGHT, before=2, after=2)
        r = p1.add_run(f"({b['num']})"); set_run_font(r, size=BODY_PT)
        for c in (c0, c1):
            tcPr = c._tc.get_or_add_tcPr(); v = OxmlElement("w:vAlign"); v.set(qn("w:val"), "center"); tcPr.append(v)
        sp = self.doc.add_paragraph(); pfmt(sp, after=0); r = sp.add_run(""); set_run_font(r, size=2)
        self.after_heading = True   # paragraph after an equation continues the sentence

    def figure(self, b, mathc, lab, cit):
        full = b["width"] == "full"
        self.mode_to(1 if full else 2)
        cm = b["cm"] or ((16.0 if full else 11.5) if REVIEW else (16.6 if full else 8.2))
        p = self.doc.add_paragraph(); pfmt(p, align=WD_ALIGN_PARAGRAPH.CENTER, before=4, after=2, keep_next=True)
        if b["label"] in BREAK_BEFORE and not REVIEW:
            p.paragraph_format.page_break_before = True
        p.add_run().add_picture(str(HERE / b["path"]), width=Cm(cm))
        c = self.doc.add_paragraph(); pfmt(c, align=WD_ALIGN_PARAGRAPH.JUSTIFY, after=8)
        r = c.add_run(f"Fig. {b['num']}. "); set_run_font(r, size=SMALL_PT, bold=True)
        add_text(c, resolve(b["caption"], lab, cit), SMALL_PT, mathc=mathc, subs=True)
        if full:
            self.mode_to(2)
        self.after_heading = True

    def table(self, b, mathc, lab, cit, width_tw=None):
        full = b["width"] == "full"
        self.mode_to(1 if full else 2)
        total = width_tw or (TEXT_W_TW if (full or REVIEW) else COL_W_TW)
        cap = self.doc.add_paragraph(); pfmt(cap, align=WD_ALIGN_PARAGRAPH.JUSTIFY, before=4, after=3, keep_next=True)
        if b["label"] in BREAK_BEFORE and not REVIEW:
            cap.paragraph_format.page_break_before = True
        r = cap.add_run(f"Table {b['num']}. "); set_run_font(r, size=SMALL_PT, bold=True)
        add_text(cap, resolve(b["caption"], lab, cit), SMALL_PT, mathc=mathc, subs=True)
        rows = b["rows"]; ncol = len(rows[0])
        wts = b["cols"] or [1] * ncol
        widths = [int(total * w / sum(wts)) for w in wts]
        tbl = self.doc.add_table(rows=len(rows), cols=ncol)
        table_setup(tbl, widths, total)
        for i, row in enumerate(rows):
            if len(row) != ncol:
                raise ValueError(f"{b['label']}: row {i} has {len(row)} cells, expected {ncol}: {row}")
            no_split_row(tbl.rows[i])
            if i == 0: header_row(tbl.rows[i])
            for j, txt in enumerate(row):
                c = tbl.rows[i].cells[j]; cell_margins(c)
                p = c.paragraphs[0]; pfmt(p, after=0, before=0, align=WD_ALIGN_PARAGRAPH.LEFT)
                add_text(p, resolve(txt, lab, cit), SMALL_PT, bold=(i == 0), mathc=mathc, subs=True)
                if i < len(rows) - 1:
                    p.paragraph_format.keep_with_next = True
                edges = {}
                if i == 0: edges.update(top=8, bottom=4)
                if i == len(rows) - 1: edges.update(bottom=8)
                if edges: cell_borders(c, **edges)
        if b.get("note"):
            n = self.doc.add_paragraph(); pfmt(n, align=WD_ALIGN_PARAGRAPH.JUSTIFY, before=2, after=6)
            add_text(n, resolve(b["note"], lab, cit), SMALL_PT - 0.5, mathc=mathc, subs=True)
        else:
            sp = self.doc.add_paragraph(); pfmt(sp, after=4)
        if full:
            self.mode_to(2)
        self.after_heading = True

    # front matter -----------------------------------------------------------
    def front(self, mathc, lab, cit):
        d = self.doc
        p = d.add_paragraph(); pfmt(p, align=WD_ALIGN_PARAGRAPH.CENTER, after=1)
        add_text(p, meta.HEADER, 12, bold=True)
        p = d.add_paragraph(); pfmt(p, align=WD_ALIGN_PARAGRAPH.CENTER, after=14)
        add_text(p, meta.HEADER2, SMALL_PT)
        p = d.add_paragraph(); pfmt(p, after=8)
        add_text(p, meta.TITLE, 16, bold=True)
        p = d.add_paragraph(); pfmt(p, after=3)
        for k, (name, aff, corr) in enumerate(meta.AUTHORS):
            if k: r = p.add_run(", "); set_run_font(r, size=11)
            r = p.add_run(name); set_run_font(r, size=11)
            r = p.add_run(aff + ("*" if corr else "")); set_run_font(r, size=11, sup=True)
        for key, txt in meta.AFFILIATIONS:
            p = d.add_paragraph(); pfmt(p, after=1)
            r = p.add_run(key); set_run_font(r, size=SMALL_PT, sup=True, italic=True)
            r = p.add_run(" " + txt); set_run_font(r, size=SMALL_PT, italic=True)
        p = d.add_paragraph(); pfmt(p, after=10)
        r = p.add_run("* "); set_run_font(r, size=SMALL_PT)
        r = p.add_run(meta.CORRESPONDING); set_run_font(r, size=SMALL_PT)
        # article info / abstract table (as in the template)
        widths = [2127, 450, TEXT_W_TW - 2127 - 450]
        t = d.add_table(rows=2, cols=3); table_setup(t, widths, TEXT_W_TW)
        heads = ["A R T I C L E   I N F O", "", "A B S T R A C T"]
        for j in range(3):
            c = t.rows[0].cells[j]; cell_margins(c, 40, 40, 0, 60)
            pp = c.paragraphs[0]; pfmt(pp, before=6, after=6)
            r = pp.add_run(heads[j]); set_run_font(r, size=9, bold=True)
            if j != 1:
                cell_borders(c, bottom=4); cell_borders(t.rows[1].cells[j], bottom=4)
        c = t.rows[1].cells[0]; cell_margins(c, 40, 40, 0, 60)
        pp = c.paragraphs[0]; pfmt(pp, after=3); r = pp.add_run("Keywords:"); set_run_font(r, size=SMALL_PT, italic=True)
        for kw in meta.KEYWORDS:
            pp = c.add_paragraph(); pfmt(pp, after=1.5); r = pp.add_run(kw); set_run_font(r, size=SMALL_PT)
        c = t.rows[1].cells[2]; cell_margins(c, 40, 80, 0, 60)
        pp = c.paragraphs[0]; pfmt(pp, align=WD_ALIGN_PARAGRAPH.JUSTIFY, after=4, first=0.4)
        add_text(pp, meta.ABSTRACT, 9, mathc=mathc)
        sp = d.add_paragraph(); pfmt(sp, after=6)
        # highlights
        p = d.add_paragraph(); pfmt(p, before=2, after=3, keep_next=True); add_text(p, "HIGHLIGHTS", 11, bold=True)
        for h in meta.HIGHLIGHTS:
            p = d.add_paragraph(); pfmt(p, after=1.5, left=0.5, hanging=0.35)
            p.paragraph_format.tab_stops.add_tab_stop(Cm(0.5))
            r = p.add_run("•\t" + h); set_run_font(r, size=9)
        self.toc()
        p = d.add_paragraph(); pfmt(p, before=8, after=3, keep_next=True); add_text(p, "NOMENCLATURE", 11, bold=True)
        outline(p, 0)
        self.nomenclature(mathc, lab, cit)
        self.after_heading = True

    def toc(self):
        d = self.doc
        p = d.add_paragraph(); pfmt(p, before=8, after=4, keep_next=True); add_text(p, "CONTENTS", 11, bold=True)
        entries = TOC_ENTRIES or [("NOMENCLATURE", 0)]
        right = Cm(TEXT_W_TW / 567.0)
        for k, (txt, lvl) in enumerate(entries):
            p = d.add_paragraph(); pfmt(p, after=1.5, left=0.6 * lvl)
            p.paragraph_format.tab_stops.add_tab_stop(right, alignment=2, leader=1)   # right, dot leader
            if k == 0:
                fld(p, "begin"); instr(p, ' TOC \\o "1-2" \\h \\z \\u '); fld(p, "separate")
            r = p.add_run(txt + "\t" + str(TOC_PAGES.get(txt, ""))); set_run_font(r, size=9, bold=(lvl == 0))
            if k == len(entries) - 1:
                fld(p, "end")

    def nomenclature(self, mathc, lab, cit):
        items = meta.NOMENCLATURE
        widths = [int(TEXT_W_TW * 0.22), TEXT_W_TW - int(TEXT_W_TW * 0.22)]
        t = self.doc.add_table(rows=len(items), cols=2); table_setup(t, widths, sum(widths))
        for i, (sym, desc) in enumerate(items):
            c0, c1 = t.rows[i].cells
            cell_margins(c0, 12, 12, 0, 60); cell_margins(c1, 12, 12, 0, 40)
            p0, p1 = c0.paragraphs[0], c1.paragraphs[0]
            pfmt(p0, after=0, before=4 if (desc is None and i) else 0); pfmt(p1, after=0)
            no_split_row(t.rows[i])
            if desc is None:
                r = p0.add_run(sym); set_run_font(r, size=SMALL_PT + 0.5, bold=True, italic=True)
                p0.paragraph_format.keep_with_next = True
            else:
                add_text(p0, sym, SMALL_PT + 0.5, mathc=mathc)
                add_text(p1, resolve(desc, lab, cit), SMALL_PT + 0.5, mathc=mathc)
        for c in t.rows[0].cells: cell_borders(c, top=8)
        for c in t.rows[-1].cells: cell_borders(c, bottom=8)
        sp = self.doc.add_paragraph(); pfmt(sp, after=6)

    def references(self, cit):
        p = self.doc.add_paragraph(); pfmt(p, before=10, after=5, keep_next=True)
        add_text(p, "REFERENCES", 11, bold=True)
        outline(p, 0)
        for k, key in enumerate(cit.order, 1):
            p = self.doc.add_paragraph()
            pfmt(p, align=WD_ALIGN_PARAGRAPH.LEFT, after=2.5, left=0.75, hanging=0.75)
            p.paragraph_format.tab_stops.add_tab_stop(Cm(0.75))
            r = p.add_run(f"[{k}]\t" + REFS[key][1]); set_run_font(r, size=SMALL_PT + 0.5)


def main():
    src = "\n\n".join((HERE / f).read_text(encoding="utf-8") for f in BODY)
    blocks = parse_source(src)
    lab = number_labels(blocks)
    cit = Citer()
    mathc = MathCache()
    # collect math everywhere
    collect_math(meta.ABSTRACT, mathc)
    for s, dsc in meta.NOMENCLATURE:
        collect_math(s, mathc)
        if dsc: collect_math(dsc, mathc)
    for b in blocks:
        if b["type"] == "eq":
            mathc.collect(b["tex"], True)
        for key in ("text", "caption", "note"):
            if b.get(key): collect_math(b[key], mathc)
        for row in b.get("rows", []):
            for c in row: collect_math(c, mathc)
    mathc.convert()
    global TOC_ENTRIES, TOC_PAGES
    TOC_ENTRIES = [("NOMENCLATURE", 0)]
    for bb in blocks:
        if bb["type"] == "h1":
            TOC_ENTRIES.append(((bb.get("num", "") + " " + bb["text"]).strip() if bb["numbered"] else bb["text"], 0))
        elif bb["type"] == "h2":
            TOC_ENTRIES.append((bb["num"] + " " + bb["text"], 1))
    TOC_ENTRIES.append(("REFERENCES", 0))
    pf = HERE / "out" / ("pages_review.json" if REVIEW else "pages_template.json")
    if pf.exists():
        import json; TOC_PAGES = json.loads(pf.read_text())
    bf = HERE / "out" / "breaks_template.json"
    if bf.exists() and not REVIEW:
        import json; BREAK_BEFORE.update(json.loads(bf.read_text()))
    B = Builder()
    B.front(mathc, lab, cit)
    B.mode_to(2)
    for b in blocks:
        t = b["type"]
        if t in ("h1", "h2", "h3"):
            B.heading(b)
        elif t == "para":
            B.para(resolve(b["text"], lab, cit), mathc)
        elif t == "item":
            b = dict(b, text=resolve(b["text"], lab, cit)); B.item(b, mathc)
        elif t == "eq":
            B.equation(b, mathc)
        elif t == "fig":
            B.figure(b, mathc, lab, cit)
        elif t == "table":
            B.table(b, mathc, lab, cit)
    B.references(cit)
    B.finish()
    core = B.doc.core_properties
    core.last_modified_by = meta.AUTHORS[0][0]; core.title = meta.TITLE; core.author = "; ".join(a[0] for a in meta.AUTHORS); core.subject = "VoxLat manuscript draft"
    core.keywords = "; ".join(meta.KEYWORDS); core.comments = "Draft generated from the VoxLat repository (Task 14)."
    OUT.parent.mkdir(exist_ok=True)
    B.doc.save(str(OUT))
    unused = [k for k in REFS if k not in cit.order]
    print("saved", OUT, "| figures", sum(b['type'] == 'fig' for b in blocks), "| tables",
          sum(b['type'] == 'table' for b in blocks), "| equations", sum(b['type'] == 'eq' for b in blocks),
          "| refs", len(cit.order), "| unused refs", unused)
    (HERE / "out" / "ref_order.txt").write_text("\n".join(f"[{i}] {k} ({REFS[k][0]})" for i, k in enumerate(cit.order, 1)))
    words = len(re.sub(r"\$[^$]*\$", "x", src).split())
    print("approx. source words:", words)


if __name__ == "__main__":
    main()
